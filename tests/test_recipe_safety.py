"""Recipe trust boundaries, bounded inputs and rendering isolation."""

from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from io import BytesIO
from threading import Barrier
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from PIL import Image
from pptx import Presentation
from recipe_cards.rendering import assets, pages
from recipe_cards.schema import MAX_IMAGES_PER_BATCH, load_recipes

RECIPE = {
    "title": "Soup",
    "slug": "soup",
    "ingredients": [{"quantity": "1 cup", "item": "Water"}],
    "steps": [{"title": "Heat", "body": "Boil the water."}],
}


def png():
    output = BytesIO()
    Image.new("RGB", (8, 8), "white").save(output, format="PNG")
    return output.getvalue()


@pytest.mark.parametrize(
    "payload",
    [
        {"recipes": "invalid"},
        {"recipes": []},
        {**RECIPE, "steps": ["not a step"]},
        {**RECIPE, "ingredients": []},
        {**RECIPE, "hero_image_path": "C:/private/image.png"},
        {**RECIPE, "hero_image_path": "https://example.com/image.png"},
        {**RECIPE, "hero_image_path": "gs://other-bucket/image.png"},
        {**RECIPE, "hero_image_path": "../../private/image.png"},
        {**RECIPE, "steps": RECIPE["steps"] * 25},
        {**RECIPE, "unexpected": "value"},
    ],
)
def test_recipe_schema_rejects_invalid_payload(payload):
    with pytest.raises(ValueError):
        load_recipes(payload)


def test_legacy_image_aliases_are_normalized():
    result = load_recipes({**RECIPE, "heroImagePath": "soup/run-1/images/hero.png"})
    assert result["recipes"][0]["hero_image_path"] == "soup/run-1/images/hero.png"


@pytest.mark.parametrize(
    "path",
    ["gs://private/secret.png", "../../private/secret.png", "/etc/passwd"],
)
def test_image_outside_the_card_store_rejected_before_download(monkeypatch, tmp_path, path):
    """A payload can only name assets inside the use-case root.

    Authorization is structural: each path is rebuilt from that root, so an
    absolute location or a traversal cannot resolve at all.
    """
    download = Mock()
    monkeypatch.setattr(assets, "download_bytes", download)
    with (
        pytest.raises(ValueError),
        assets.asset_context({"hero_image_path": path}, "test", str(tmp_path), "test-bucket"),
    ):
        pytest.fail("Unauthorized image accepted")
    download.assert_not_called()


def test_concurrent_asset_contexts_remain_isolated(monkeypatch, tmp_path):
    monkeypatch.setattr(assets, "download_bytes", lambda *args, **kwargs: png())
    barrier = Barrier(2)

    def render(index):
        directory = tmp_path / str(index)
        directory.mkdir()
        uri = f"soup/run-{index}/images/hero.png"
        with assets.asset_context({"hero_image_path": uri}, "test", str(directory), "test-bucket"):
            before = assets.resolve(uri)
            barrier.wait(timeout=10)
            assert before == assets.resolve(uri)
            assert assets.resolve(f"soup/run-{1 - index}/images/hero.png") == ""
        assert assets.CURRENT_ASSETS.get() is None

    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(render, (0, 1)))


def test_failed_download_resets_asset_context(monkeypatch, tmp_path):
    monkeypatch.setattr(assets, "download_bytes", Mock(side_effect=ValueError("too large")))
    uri = "soup/run-1/images/hero.png"
    with (
        pytest.raises(ValueError, match="too large"),
        assets.asset_context({"hero_image_path": uri}, "test", str(tmp_path), "test-bucket"),
    ):
        pytest.fail("Download unexpectedly succeeded")
    assert assets.CURRENT_ASSETS.get() is None


def test_render_keeps_complete_instruction():
    deck = Presentation(BytesIO(pages.render_deck(RECIPE, "test")))
    text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "Boil the water." in text


@pytest.mark.parametrize("core_count,extra_count", [(12, 1), (13, 0), (24, 12), (5, 12)])
def test_all_ingredients_and_quantities_are_rendered(core_count, extra_count):
    recipe = deepcopy(RECIPE)
    recipe["ingredients"] = [
        {"quantity": f"{i + 1} g", "item": f"Core {i + 1}"} for i in range(core_count)
    ]
    recipe["variation_ingredients"] = [
        {"quantity": f"{i + 25} g", "item": f"Extra {i + 1}"} for i in range(extra_count)
    ]
    deck = Presentation(BytesIO(pages.render_deck(recipe, "test")))
    texts = [shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]
    for ingredient in recipe["ingredients"] + recipe["variation_ingredients"]:
        assert texts.count(ingredient["item"]) == 1
        assert texts.count(ingredient["quantity"]) == 1
    assert not any(text.startswith("+ ") and text.endswith("more") for text in texts)
    if extra_count:
        assert "OPTIONAL VARIATION" in texts


def test_overlong_ingredient_quantity_returns_a_correction():
    from recipe_cards.errors import ContentTooLong

    recipe = deepcopy(RECIPE)
    recipe["ingredients"][0]["quantity"] = "one hundred and twenty five grams of water"
    with pytest.raises(ContentTooLong, match="Ingredient quantity") as raised:
        pages.render_deck(recipe, "test")
    assert raised.value.field == "ingredients[0].quantity"


def test_overflow_raises_instead_of_truncating():
    """Cooking text is never silently shortened.

    The renderer stops and says which step overflowed and by how much, so the
    caller can shorten it. Dropping the words instead would hand someone a card
    missing part of a step, with nothing to show it had happened.
    """
    from recipe_cards.errors import ContentTooLong

    recipe = deepcopy(RECIPE)
    recipe["steps"][0]["body"] = " ".join(["Keep stirring the soup carefully."] * 30)
    with pytest.raises(ContentTooLong) as raised:
        pages.render_deck(recipe, "test")

    # The failure names the field to edit and the edit to make.
    assert raised.value.field.startswith("steps[")
    assert "words" in raised.value.suggestion


def test_image_batch_limits_before_paid_calls(monkeypatch):
    from recipe_cards import images

    monkeypatch.setattr(images, "OUTPUT_BUCKET", "output")
    generate = Mock()
    monkeypatch.setattr(images, "generate_images", generate)
    with pytest.raises(ValueError, match="At most"):
        images.generate_recipe_images(
            "soup",
            ["photo"] * (MAX_IMAGES_PER_BATCH + 1),
            [str(i) for i in range(MAX_IMAGES_PER_BATCH + 1)],
            SimpleNamespace(state={}),
        )
    generate.assert_not_called()


def test_image_names_cannot_collide_after_normalization(monkeypatch):
    from recipe_cards import images

    monkeypatch.setattr(images, "OUTPUT_BUCKET", "output")
    with pytest.raises(ValueError, match="unique after normalization"):
        images.generate_recipe_images(
            "soup", ["photo", "photo"], ["Step 1", "step-1"], SimpleNamespace(state={})
        )


@pytest.mark.parametrize("use_references", [True, False])
def test_reference_free_sketch_omits_style_plates(monkeypatch, use_references):
    from recipe_cards import images

    monkeypatch.setattr(images, "OUTPUT_BUCKET", "output")
    plates = Mock(return_value=(b"style-photo",))
    monkeypatch.setattr(images, "_style_plates", plates)
    generate = Mock(
        return_value=[SimpleNamespace(name="sketch", data=png(), mime_type="image/png")]
    )
    monkeypatch.setattr(images, "generate_images", generate)
    monkeypatch.setattr(images, "upload_bytes", Mock())
    images.generate_recipe_images(
        "soup",
        ["Navy ink sketch"],
        ["sketch"],
        SimpleNamespace(state={}),
        use_reference_images=use_references,
    )
    request = generate.call_args.args[0][0]
    assert request.reference_images == ((b"style-photo",) if use_references else ())
    assert generate.call_args.kwargs["mode"] == "parallel"
    assert plates.call_count == int(use_references)


def test_no_references_cannot_enable_photo_chaining(monkeypatch):
    from recipe_cards import images

    monkeypatch.setattr(images, "OUTPUT_BUCKET", "output")
    generate = Mock()
    monkeypatch.setattr(images, "generate_images", generate)
    with pytest.raises(ValueError, match="No-reference"):
        images.generate_recipe_images(
            "soup",
            ["sketch"],
            ["sketch"],
            SimpleNamespace(state={}),
            mode="sequential_reference",
            use_reference_images=False,
        )
    generate.assert_not_called()


def test_generated_assets_are_owned_by_session(monkeypatch):
    from recipe_cards import images, publish

    monkeypatch.setattr(images, "OUTPUT_BUCKET", "output")
    monkeypatch.setattr(publish, "OUTPUT_BUCKET", "output")
    monkeypatch.setattr(
        images,
        "generate_images",
        lambda *args, **kwargs: [SimpleNamespace(name="hero", data=png(), mime_type="image/png")],
    )
    monkeypatch.setattr(
        images,
        "upload_bytes",
        lambda project, bucket, name, *args, **kwargs: f"gs://{bucket}/{name}",
    )
    context = SimpleNamespace(state={})
    result = images.generate_recipe_images("soup", ["photo"], ["hero"], context)
    recipe = {**RECIPE, "hero_image_path": result["images"]["hero"]}
    import json

    with pytest.raises(ValueError, match="Unknown recipe run"):
        publish.render_recipe_card(json.dumps(recipe), SimpleNamespace(state={}), result["run_id"])
    render = Mock(return_value=b"deck")
    monkeypatch.setattr(publish, "render_deck", render)
    monkeypatch.setattr(publish, "upload_bytes", Mock(return_value="gs://output/deck.pptx"))
    publish.render_recipe_card(json.dumps(recipe), context, result["run_id"])
    # Generation hands back a relative path, and that is what the payload
    # carries: the model never sees or writes a bucket name.
    hero = result["images"]["hero"]
    assert not hero.startswith("gs://")
    assert render.call_args.args[0]["recipes"][0]["hero_image_path"] == hero
    # The bucket reaches the renderer as configuration, not through the payload.
    assert render.call_args.kwargs["bucket"] == publish.OUTPUT_BUCKET


def test_bounded_download_reads_only_limit_plus_one(monkeypatch):
    from gemini_shared.connectors import cloud_storage

    client = Mock()
    blob = client.bucket.return_value.blob.return_value
    blob.download_as_bytes.return_value = b"12345"
    monkeypatch.setattr(cloud_storage.storage, "Client", Mock(return_value=client))
    with pytest.raises(ValueError, match="download limit"):
        cloud_storage.download_bytes("test", "gs://output/large.png", max_bytes=4)
    blob.download_as_bytes.assert_called_once_with(start=0, end=4, raw_download=True)


def test_authorized_image_is_embedded_in_deck(monkeypatch):
    monkeypatch.setattr(assets, "download_bytes", lambda *args, **kwargs: png())
    uri = "soup/run/images/hero.png"
    recipe = {**RECIPE, "hero_image_path": uri, "decorative_image_path": uri}
    deck = Presentation(BytesIO(pages.render_deck(recipe, "test", bucket="test-bucket")))
    assert any(shape.shape_type == 13 for slide in deck.slides for shape in slide.shapes)
    assert assets.CURRENT_ASSETS.get() is None


def test_workflow_validation_blocks_bad_recipe_before_images():
    from recipe_card_workflow.stages.validation import validate_recipe_stage

    context = SimpleNamespace(state={"recipe": '{"recipes":"invalid"}'})
    with pytest.raises(ValueError):
        validate_recipe_stage(context)


@pytest.mark.parametrize("state", [{}, {"recipe": ""}, {"recipe": "   "}])
def test_workflow_validation_reports_a_missing_recipe(state):
    """An empty recipe means the writing stage produced nothing.

    That reached load_recipes as "" and raised a JSONDecodeError naming this
    file, which pointed at the wrong stage entirely.
    """
    from recipe_card_workflow.stages.validation import validate_recipe_stage

    with pytest.raises(ValueError, match="recipe_writer"):
        validate_recipe_stage(SimpleNamespace(state=state))


def test_workflow_validation_normalizes_recipe():
    import json

    from recipe_card_workflow.stages.validation import validate_recipe_stage

    context = SimpleNamespace(state={"recipe": json.dumps(RECIPE)})
    validate_recipe_stage(context)
    assert json.loads(context.state["recipe"])["title"] == "Soup"


def test_image_client_is_closed_after_attempt(monkeypatch):
    from gemini_shared.media import images as media

    client = Mock()
    manager = Mock()
    manager.__enter__ = Mock(return_value=client)
    manager.__exit__ = Mock(return_value=False)
    monkeypatch.setattr(media, "_client", Mock(return_value=manager))
    monkeypatch.setattr(media, "_wait_for_slot", lambda: None)
    monkeypatch.setattr(media, "_extract_image", lambda *args: (png(), "image/png"))
    result = media._generate_one("global", "test-model", media.ImageRequest("photo", "hero"), ())
    assert result.name == "hero"
    manager.__exit__.assert_called_once_with(None, None, None)


def test_a_short_step_stays_shorter_than_the_step_beside_it():
    """Levelling columns must not invert a row's measured difference.

    Each block is sized from its own text, so two steps in a row differ. Handing
    a short column's whole leftover height to its last block made that block
    taller than the longer one beside it, so every row looked the same height
    whatever it carried. The block height is asserted rather than the
    photograph's, because a step without an image draws a placeholder and there
    is no picture to measure.
    """
    from recipe_cards.rendering.steps import place_step_blocks

    long_step = {"title": "Longer one", "bullets": ["Stir the pot gently for a minute."] * 4}
    short_step = {"title": "Short", "bullets": ["Brief."]}
    steps = [
        {"title": "First", "bullets": ["Stir the pot gently for a minute."] * 2},
        {"title": "Second", "bullets": ["Stir the pot gently for a minute."] * 2},
        long_step,
        short_step,
    ]

    placed = place_step_blocks(steps, top=1.30, bottom=10.90)
    bottom_row = [placed[2][3], placed[3][3]]

    assert bottom_row[0] > bottom_row[1], (
        "the shorter step's block is at least as tall as the longer step's"
    )


@pytest.mark.parametrize("state", [{}, {"images": ""}, {"images": "   "}])
def test_workflow_validation_reports_missing_images(state):
    """An absent mapping killed the run with no usable error.

    card_renderer interpolates {images} from state, so a missing key raised a
    KeyError while the instruction was built, before any tool ran. Gemini
    Enterprise reported only FAILED_PRECONDITION with empty error details,
    naming neither the stage nor the cause.
    """
    from recipe_card_workflow.stages.validation import validate_images_stage

    with pytest.raises(ValueError, match="image_director"):
        validate_images_stage(SimpleNamespace(state=state))


def test_workflow_validation_rejects_an_empty_image_mapping():
    """A mapping with no paths would render a deck of placeholders."""
    import json

    from recipe_card_workflow.stages.validation import validate_images_stage

    context = SimpleNamespace(state={"images": json.dumps({"run_id": "r", "images": {}})})
    with pytest.raises(ValueError, match="no image paths"):
        validate_images_stage(context)


def test_workflow_validation_reports_unparsable_images():
    from recipe_card_workflow.stages.validation import validate_images_stage

    context = SimpleNamespace(state={"images": "Here are your images!"})
    with pytest.raises(ValueError, match="JSON object"):
        validate_images_stage(context)


def test_workflow_validation_accepts_a_real_image_mapping():
    import json

    from recipe_card_workflow.stages.validation import validate_images_stage

    mapping = {"run_id": "20260912-200323-abc", "images": {"hero": "dish/run/images/hero.png"}}
    validate_images_stage(SimpleNamespace(state={"images": json.dumps(mapping)}))


def test_the_renderer_stage_guards_its_input():
    """The guard only helps if it is actually wired onto the stage."""
    from recipe_card_workflow.stages import card_renderer
    from recipe_card_workflow.stages.validation import validate_images_stage

    assert card_renderer.before_agent_callback is validate_images_stage
