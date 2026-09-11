"""Recipe substitutions preserve quantities, instructions and editable accents."""

from copy import deepcopy
from io import BytesIO

import pytest
from pptx import Presentation
from recipe_cards.rendering.pages import render_deck
from recipe_cards.schema import load_recipes

RECIPE = {
    "title": "Pork & Chive Potstickers",
    "slug": "pork-chive-potstickers",
    "ingredients": [
        {"item": "Ground pork", "quantity": "1 lb"},
        {"item": "Chives", "quantity": "2 cups"},
        {"item": "Wrappers", "quantity": "30"},
    ],
    "steps": [{"title": "Mix the filling", "body": "Mix the pork and chives."}],
    "customizations": [
        {
            "name": "Shrimp & Pork",
            "replaces": ["Ground pork"],
            "ingredients": [
                {"item": "Ground pork", "quantity": "8 oz"},
                {"item": "Peeled shrimp", "quantity": "8 oz"},
            ],
            "steps": [
                {
                    "step": 1,
                    "instructions": [
                        "Finely mince the shrimp.",
                        "Mix the shrimp with 8 oz pork and the remaining filling ingredients.",
                    ],
                }
            ],
        }
    ],
}


def slide_text(slide):
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


def test_customized_quantities_and_step_checklists_render():
    deck = Presentation(BytesIO(render_deck(RECIPE, "test")))
    first = slide_text(deck.slides[0])
    assert "CUSTOMIZED OPTIONS" in first
    assert "BASE OPTION" in first
    assert "1 lb" in first and first.count("8 oz") == 2
    assert "Peeled shrimp" in first and "SHRIMP & PORK" in first
    all_text = "\n".join(slide_text(slide) for slide in deck.slides)
    assert "Customized Steps" in all_text and "STEP 1" in all_text
    assert "Finely mince the shrimp." in all_text
    assert "See Customized Steps" in all_text
    accents = [s for s in deck.slides[0].shapes if s.name == "Menu banner accent"]
    assert len(accents) == 2
    assert all(s._element.xpath("./p:spPr/a:custGeom") for s in accents)
    assert any(s._element.xpath("./p:spPr/a:gradFill") for s in accents)


@pytest.mark.parametrize("bad", ["step", "ingredient"])
def test_customizations_reference_existing_recipe_parts(bad):
    recipe = deepcopy(RECIPE)
    if bad == "step":
        recipe["customizations"][0]["steps"][0]["step"] = 2
    else:
        recipe["customizations"][0]["replaces"] = ["Chicken"]
    with pytest.raises(ValueError):
        load_recipes(recipe)


def test_long_customized_checklists_continue_without_losing_text():
    recipe = deepcopy(RECIPE)
    instructions = [f"Instruction {i}: " + "Mix gently before proceeding. " * 4 for i in range(12)]
    recipe["customizations"][0]["steps"] = [
        {"step": 1, "instructions": instructions} for _ in range(3)
    ]
    deck = Presentation(BytesIO(render_deck(recipe, "test")))
    text = "\n".join(slide_text(slide) for slide in deck.slides)
    for instruction in instructions:
        assert text.count(instruction) == 3
    assert text.count("Customized Steps") >= 2


def test_banner_asset_is_in_both_deployment_packages():
    from deploy.sources import source_entries
    from registry import get_agent_spec

    for name in ("recipe_card_agent", "recipe_card_workflow"):
        names = {entry.archive_name for entry in source_entries(get_agent_spec(name))}
        assert "recipe_cards/rendering/banner_shapes.json" in names
