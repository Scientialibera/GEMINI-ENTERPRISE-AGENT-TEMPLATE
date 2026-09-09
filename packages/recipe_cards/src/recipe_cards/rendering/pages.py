"""Assemble recipe decks from independently maintained page layouts."""

from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches

from ..schema import load_recipes
from .assets import asset_context
from .drawing import (
    chunks,
)
from .ingredients import add_ingredient_continuation_slides
from .overview import add_overview_slide
from .steps import add_steps_slide
from .theme import (
    PAGE_H,
    PAGE_W,
)


def add_recipe(prs, recipe):
    add_overview_slide(prs, recipe)
    ingredient_pages = add_ingredient_continuation_slides(prs, recipe)
    steps = list(recipe.get("steps") or [])
    step_pages = list(chunks(steps, 4)) or [[]]
    for page_index, group in enumerate(step_pages):
        add_steps_slide(
            prs,
            recipe,
            page_index,
            page_index * 4,
            group,
            len(step_pages),
            page_offset=1 + ingredient_pages,
        )


def build_pptx(data: dict[str, Any], output: str) -> None:
    """Write the deck for every recipe in ``data`` to ``output``."""
    prs = Presentation()
    prs.slide_width = Inches(PAGE_W)
    prs.slide_height = Inches(PAGE_H)

    # Remove the default first slide if present (normally Presentation() has none).
    while len(prs.slides) > 0:
        r_id = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[0]

    recipes = data.get("recipes") or []
    if not recipes:
        raise ValueError("Input contains no recipes.")

    for recipe in recipes:
        add_recipe(prs, recipe)

    Path(output).parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def render_deck(
    data: dict[str, Any], project_id: str, *, allowed_uris: set[str] | None = None
) -> bytes:
    """Render validated content using only the run's authorized images."""
    data = load_recipes(data)
    with (
        tempfile.TemporaryDirectory(prefix="recipe-card-") as work_dir,
        asset_context(data, project_id, work_dir, allowed_uris or set()),
    ):
        output = Path(work_dir) / "recipe_cards.pptx"
        build_pptx(data, str(output))
        return output.read_bytes()
